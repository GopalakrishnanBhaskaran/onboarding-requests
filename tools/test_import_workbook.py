#!/usr/bin/env python3
"""Tests for tools/import_workbook.py on a SAMPLE workbook in the layout of the real ones (no real data).

  python3 tools/test_import_workbook.py            # run the checks
  python3 tools/test_import_workbook.py --fixture  # also rewrite tests/fixtures/import-sample.json (read by the .NET tests)
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import import_workbook as iw  # noqa: E402

RED, AMBER = "FFFCE4E4", "FFFFF2CC"
HEADER = ["Tenant", "Environment", "Group Name", "Object ID (if available)", "Role / Access Type", "Intended Job Role(s)", "Scope", "Purpose", "Source File",
          "Notes / Ambiguities"]


def workbook(folder: Path) -> tuple[Path, Path]:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    c = Workbook()
    c.remove(c.active)
    tabs = {
        "PEST - Dev": [
            ("azr-pes-rg-data-dev-ph-weu-support", "11111111-1111-4111-8111-111111111111", "CONTRIBUTOR - whole platform\n\nAzure roles held:",
             "Domain: Central hub / platform - THIS IS THE dai DOMAIN (ph)\nResources this group holds a role on:\n  • Storage Blob Data Contributor → samplestddevdaiweu01\n  • Reader → pes-rg-data-dev-dai-weu-001 / pes-rg-data-dev-acq-weu-001",
             "Application / Production Support", "• BROAD ACCESS: whole platform", AMBER),
            ("azr-pes-rg-data-dev-fin-weu-powerbidevelopers", "22222222-2222-4222-8222-222222222222", "REPORTING - no data-plane RBAC",
             "Domain: Finance (fin)\nResources: NONE DOCUMENTED.", "Power BI Developer", "• Members: 2", None),
            ("azr-pes-rg-data-dev-sup-weu-powerbideveloper", "not captured", "TYPO", "Domain: Supply  [PEST spelling of supp] (sup)", "", "• EMPTY GROUP", RED),
        ],
        "PEST - Prod": [],
        "Vestacy - Dev": [
            ("azr-ves-rg-data-dev-ph-weu-dataengineers", "33333333-3333-4333-8333-333333333333", "CONTRIBUTOR", "Domain: Central hub (ph)", "Data Engineer", "", None),
        ],
    }
    for title, rows in tabs.items():
        ws = c.create_sheet(title)
        ws.append([f"{title} (SAMPLE)"])
        ws.append(["SAMPLE layout of the classification workbook."])
        ws.append(HEADER)
        for name, oid, role, scope, jobs, notes, fill in rows:
            ws.append([title.split(" - ")[0], title.split(" - ")[1], name, oid, role, jobs, scope, "", "SAMPLE.csv", notes])
            if fill:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill("solid", fgColor=fill)
        if not rows:
            ws.append(["NO GROUPS EXIST - see the note above."])
    cpath = folder / "classification-sample.xlsx"
    c.save(cpath)

    r = Workbook()
    ws = r.active
    ws.title = "RBAC"
    ws.append(["", "MI-ResourceName", "IdentityType", "RoleName", "ResourceName", "ResourceType", "Notes"])
    ws.append([1, "azr-pes-rg-data-dev-ph-weu-support-reader", "Entra ID Group", "Storage Blob Data Reader", "samplestddevdaiweu01\nsamplestddevlogweu01",
               "Microsoft.Storage/storageAccounts", "SAMPLE support reader"])
    ws.append(["", "", "", "Reader", "pes-rg-data-dev-dai-weu-001", "Microsoft.Resources/resourceGroups", ""])
    ws.append([2, "azr-pes-rg-data-prd-ph-weu-support-reader", "Entra ID Group", "Reader", "pes-rg-data-prd-dai-weu-001", "Microsoft.Resources/resourceGroups", ""])
    rpath = folder / "response-sample.xlsx"
    r.save(rpath)
    return cpath, rpath


def run(write_fixture: bool) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        cpath, rpath = workbook(Path(tmp))
        scope = "/subscriptions/00000000-0000-0000-0000-000000000000/resourceGroups/pes-rg-data-dev-dai-weu-001"
        data = iw.build("PEST", "SAMPLE-PEST-IMPORT", cpath, [rpath], {"pes-rg-data-dev-dai-weu-001": scope}, use_az=False, resource_prefix="SAMPLE ")
    rows = {r["entraGroup"]: r for r in data["mappingSheet"]}

    # Tenant separation: the Vestacy tab is never read for PEST.
    assert not any("-ves-" in g for g in rows), rows.keys()
    assert data["project"]["environments"] == ["Dev", "Prod"], data["project"]

    support = rows["azr-pes-rg-data-dev-ph-weu-support"]
    assert (support["environment"], support["accessLevel"], support["flag"], support["groupState"]) == ("Dev", "Support", "Amber", "Exists")
    assert support["resource"] == "SAMPLE PEST Central hub / platform (ph)", support["resource"]
    assert support["groupObjectId"] == "11111111-1111-4111-8111-111111111111"
    assert [(g["roleName"], g["resourceName"]) for g in support["azureRoles"]] == [
        ("Storage Blob Data Contributor", "samplestddevdaiweu01"), ("Reader", "pes-rg-data-dev-dai-weu-001"), ("Reader", "pes-rg-data-dev-acq-weu-001")]
    assert [g["scope"] for g in support["azureRoles"]] == [None, scope, None]

    # Privileged exactly as the review reads the role ("CONTRIBUTOR - whole platform"): flagged, so requests are not returned for it.
    assert support["privileged"] is True and rows["azr-pes-rg-data-dev-fin-weu-powerbidevelopers"]["privileged"] is False
    assert iw.has_word("CONTRIBUTOR - data read/write", iw.PRIVILEGED_WORDS) and not iw.has_word("Support Reader", iw.PRIVILEGED_WORDS)
    assert iw.has_word("Full Control of the vault", iw.PRIVILEGED_WORDS) and not iw.has_word("Storage Blob Data Contributors list", ["contributor"])
    assert rows["azr-pes-rg-data-dev-fin-weu-powerbidevelopers"]["azureRoles"] == []
    typo = rows["azr-pes-rg-data-dev-sup-weu-powerbideveloper"]
    assert (typo["flag"], typo["groupObjectId"], typo["resource"]) == ("Red", None, "SAMPLE PEST Supply (sup)")

    reader = rows["azr-pes-rg-data-dev-ph-weu-support-reader"]
    assert (reader["groupState"], reader["flag"], reader["accessLevel"]) == ("ToBeCreated", "Amber", "Support Reader")
    assert len(reader["azureRoles"]) == 3 and reader["azureRoles"][2] == {
        "roleName": "Reader", "resourceName": "pes-rg-data-dev-dai-weu-001", "resourceType": "Microsoft.Resources/resourceGroups", "scope": scope, "roleDefinitionId": None}

    prod = rows["azr-pes-rg-data-prd-ph-weu-support-reader"]
    assert (prod["environment"], prod["flag"], prod["groupState"]) == ("Prod", "Red", "ToBeCreated"), prod
    assert "no Prod groups" in prod["notes"]

    # Pilot test scope: every role lands in the one test resource group, repeats of a role merged (names kept).
    with tempfile.TemporaryDirectory() as tmp:
        cpath, rpath = workbook(Path(tmp))
        test_rg = "/subscriptions/00000000-0000-0000-0000-000000000000/resourceGroups/rg-sample-test"
        tested = iw.build("PEST", "SAMPLE-PEST-IMPORT", cpath, [rpath], {}, use_az=False, test_scope=test_rg)
    t = {r["entraGroup"]: r for r in tested["mappingSheet"]}["azr-pes-rg-data-dev-ph-weu-support-reader"]
    assert [(g["roleName"], g["scope"]) for g in t["azureRoles"]] == [("Storage Blob Data Reader", test_rg), ("Reader", test_rg)], t["azureRoles"]
    assert t["azureRoles"][0]["resourceName"] == "samplestddevdaiweu01, samplestddevlogweu01", t["azureRoles"][0]
    assert t["notes"].startswith("TEST: the roles are assigned in the test resource group rg-sample-test"), t["notes"]

    # Domain labels: notes in brackets or parentheses are dropped, the code stays.
    assert iw.clean_domain("Data Governance (domain, not role suffix)", "datagovernance") == "Data Governance (datagovernance)"
    assert iw.clean_domain("Supply  [PEST spelling of supp] (sup)", "sup") == "Supply (sup)"

    # The project document: real (not sample), tenant, domain, one test contact everywhere; live values kept otherwise.
    doc = iw.project_doc(data, None, "", "@Sample.invalid", "00000000-0000-4000-8000-0000000000AA", "Tester@Sample.invalid")
    assert (doc["name"], doc["tenantName"], doc["isSample"], doc["domainIdSuffix"], doc["tenantId"]) == (
        "PEST data platform", "PEST", False, "@sample.invalid", "00000000-0000-4000-8000-0000000000aa"), doc
    assert doc["contacts"]["security"] == ["tester@sample.invalid"] and doc["contacts"]["idTeamManager"] == "tester@sample.invalid"
    kept = iw.project_doc(data, {**doc, "name": "PEST (live name)"}, "", None, None, None)
    assert (kept["name"], kept["contacts"], kept["tenantId"]) == ("PEST (live name)", doc["contacts"], doc["tenantId"])
    assert iw.users(["sample-user-1=Requester,Owner"], "PEST") == [
        {"id": "sample-user-1:PEST", "userId": "sample-user-1", "projectId": "PEST", "roles": ["Requester", "Owner"], "isSample": False, "schemaVersion": 1}]
    # By sign-in name: stored as "signin:<lower case mail>" (the API matches guest and live.com# forms of it).
    assert iw.users(["Sample.Tester@Sample.invalid=Security"], "PEST") == [
        {"id": "signin:sample.tester@sample.invalid:PEST", "userId": "signin:sample.tester@sample.invalid", "projectId": "PEST", "roles": ["Security"],
         "isSample": False, "schemaVersion": 1}]
    for bad in (["sample-user-1=Approver"], ["bad id!=Requester"], ["sample-user-1="], ["live.com#a@sample.invalid=Owner"],
                ["a_sample.invalid#EXT#@t.invalid=Owner"], ["signin:x=Owner"], ["not a mail@x=Owner"]):
        try:
            iw.users(bad, "PEST")
            raise AssertionError(f"accepted {bad}")
        except SystemExit:
            pass

    for r in data["mappingSheet"]:  # what MappingSheetService would accept
        assert iw.re.match(r"^[A-Za-z0-9_.:-]{1,64}$", r["rowId"]) and r["resource"] and r["accessLevel"] and r["role"] and r["sourceReference"]

    if write_fixture:
        out = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "import-sample.json"
        for r in data["mappingSheet"]:
            r["updatedAt"] = "2026-10-21T00:00:00+00:00"
            r["isSample"] = True
            r["sourceReference"] = "SAMPLE workbook (tools/test_import_workbook.py)"
        out.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print("wrote", out)
    print("import_workbook: all checks passed")


if __name__ == "__main__":
    run("--fixture" in sys.argv)
