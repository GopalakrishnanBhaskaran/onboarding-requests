#!/usr/bin/env python3
"""Import the Entra group classification workbook and Data Response documents into a project's mapping sheet.

One project per Entra tenant: only the tabs of --tenant are read ("<Tenant> - <Environment>"), so Vestacy and PEST are
never mixed. Real workbook data stays out of git: the output goes to .local/ (git-ignored) and, with --apply, into Cosmos.

Inputs
  --classification  Entra-Group-Classification-*.xlsx: one row per live group (Group Name, Object ID, Role / Access Type,
                    Intended Job Role(s), Scope, Notes). Row colour: red = wrong or undecided (the AI review never passes
                    it), amber = usable with a caveat (shown to the approvers).
  --response        Data Response document(s), sheet RBAC: MI-ResourceName (the group), RoleName, ResourceName (one per
                    line), ResourceType, Notes; continuation rows without a group belong to the group above. Every role x
                    resource becomes one Azure role of the group.
  --scope-map       JSON {"resource name": "/subscriptions/..."} to resolve resource names to Azure scopes. Names not in
                    the map stay unresolved: the role assignment team assigns those (the agent assigns a group's roles
                    only when every one is resolved and allow-listed).
  --resolve         Resolve the remaining names with the Azure CLI (az resource list / az group list; read-only).
  --test-scope      Pilot test in a tenant without the real resources: every Azure role of every row is given this one
                    resource group scope (the real resource names stay on each line, so approvers see what it stands
                    for). Roles repeated by that are merged. Never for production.

With --apply the project (one per tenant) is created or updated in Cosmos: tenant name, environments, group naming
pattern, and optionally --domain-suffix, --tenant-id, --test-contact (every project contact = one tester mailbox) and
--register-user USER=Role,Role (repeatable; USER = the person's sign-in name, or the userId from /.auth/me).

Groups the response document names but the classification does not list are rows "to be created" (amber); in an
environment the classification says has no groups (e.g. Prod) they are red.

Examples
  python3 tools/import_workbook.py --tenant PEST --project PEST --project-name "PEST data platform" \\
      --classification Entra-Group-Classification.xlsx --response "Data Response document - PEST.xlsx" --resolve
  COSMOS_CONNECTION_STRING=... python3 tools/import_workbook.py ... --apply [--update]
  # pilot test in the agent's own tenant: roles land in the approved test resource group, every mail to the tester
  COSMOS_CONNECTION_STRING=... python3 tools/import_workbook.py --tenant PEST --project PEST ... \\
      --test-scope /subscriptions/<sub>/resourceGroups/rg-pmo-sample-access --tenant-id <agent tenant guid> \\
      --domain-suffix @contoso.onmicrosoft.com --test-contact tester@example.com --register-user tester@example.com=Requester,Owner --apply
"""
from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ENVIRONMENTS = {"dev": "Dev", "tst": "Test", "test": "Test", "ppd": "PreProd", "prd": "Prod"}
GROUP = re.compile(r"^azr-(?P<t>[a-z]+)-rg-data-(?P<tier>dev|tst|test|ppd|prd)-(?P<domain>[a-z0-9]+)-weu-(?P<role>[A-Za-z0-9-]+)$")
GUID = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")
SCOPE = re.compile(r"^/subscriptions/[0-9a-fA-F-]{36}(/resourceGroups/[A-Za-z0-9._()-]{1,90}(/providers/[A-Za-z0-9./_-]+)?)?$")
BULLET = re.compile(r"^\s*[•\-*]\s*(?P<role>[^→]+?)\s*→\s*(?P<res>.+?)\s*$")
LEVELS = {
    "admins": "Admin", "dataengineers": "Data Engineer", "devops": "DevOps", "developers": "Developer",
    "powerbidevelopers": "Power BI Developer", "powerbideveloper": "Power BI Developer (singular duplicate)",
    "powerbippdelopers": "Power BI Developer (misspelt group)", "powerbiuser": "Power BI User", "powerbiuserpii": "Power BI User (PII)",
    "analysts": "Analyst", "datagovernance": "Data Governance", "support": "Support", "support-reader": "Support Reader",
    "support-writer": "Support Writer",
}
# Same words and matching as the AI review (ReviewPolicy.PrivilegedWords, ReviewChecks.HasWord): a row whose role reads
# as privileged must be flagged, or the review returns every request for it ("PMO must flag it").
PRIVILEGED_WORDS = ["owner", "admin", "administrator", "contributor", "user access administrator", "global", "privileged", "elevated", "full control", "manage"]
RESOURCE_TYPES = {"rg": "Microsoft.Resources/resourceGroups", "st": "Microsoft.Storage/storageAccounts", "kv": "Microsoft.KeyVault/vaults"}


def flag_of(cell) -> str:
    """Red / Amber / None from the row's fill colour (light red FCE4E4, light amber FFF2CC in the workbook)."""
    fill = getattr(cell, "fill", None)
    rgb = getattr(getattr(fill, "fgColor", None), "rgb", None) if fill is not None and fill.fill_type else None
    if not isinstance(rgb, str) or len(rgb) < 6:
        return "None"
    r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    if r > g + 12 and abs(g - b) < 24:
        return "Red"
    if r > b + 24 and g > b + 24:
        return "Amber"
    return "None"


def clean_domain(label: str, code: str) -> str:
    """'Central hub / platform - THIS IS THE dai DOMAIN … (ph)' -> 'Central hub / platform (ph)'; '[alias]' and
    '(a note, with words)' notes dropped."""
    base = re.sub(r"\s*\[[^\]]*\]", "", label).split(" - ")[0]
    base = re.sub(r"\s*\([^)]*[ ,][^)]*\)", "", base)
    base = re.sub(r"\s*\([a-z0-9]+\)\s*$", "", base).strip()
    return f"{base} ({code})"


def text(v) -> str:
    return "" if v is None else str(v).strip()


def has_word(text: str, words: list[str]) -> bool:
    """ReviewChecks.HasWord: whole words of the lower-cased text (letters, digits, dashes), or a multi-word phrase."""
    if not text or not text.strip():
        return False
    t = " " + re.sub(r"[^a-z0-9-]+", " ", text.lower()) + " "
    return any(f" {w} " in t or (" " in w and w in t) for w in words)


def level_of(role: str) -> str:
    return LEVELS.get(role.lower(), " ".join(w.capitalize() for w in re.split(r"[-_]", role) if w))


def guess_type(name: str) -> str | None:
    n = name.lower()
    if "-rg-" in n or n.startswith("rg-"):
        return RESOURCE_TYPES["rg"]
    if "-kv-" in n:
        return RESOURCE_TYPES["kv"]
    if re.match(r"^[a-z0-9]{3,24}$", n) and "std" in n:
        return RESOURCE_TYPES["st"]
    return None


def read_response(paths: list[Path]) -> dict[str, list[dict]]:
    """Group name (lower case) -> [{roleName, resourceName, resourceType, notes}] from the RBAC sheets."""
    import openpyxl
    out: dict[str, list[dict]] = {}
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["RBAC"] if "RBAC" in wb.sheetnames else wb.worksheets[0]
        header = [text(c).lower() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = {name: header.index(name) for name in ("mi-resourcename", "rolename", "resourcename", "resourcetype", "notes") if name in header}
        if not {"mi-resourcename", "rolename", "resourcename"} <= col.keys():
            raise SystemExit(f"{path}: RBAC sheet needs MI-ResourceName, RoleName and ResourceName columns")
        current = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            group = text(row[col["mi-resourcename"]])
            if group:
                current = group.lower()
            role = text(row[col["rolename"]])
            if not current or not role:
                continue
            for res in [x.strip() for x in text(row[col["resourcename"]]).splitlines() if x.strip()]:
                out.setdefault(current, []).append({
                    "roleName": role, "resourceName": res,
                    "resourceType": text(row[col["resourcetype"]]) if "resourcetype" in col else guess_type(res),
                })
    return out


def resolve(names: set[str], scope_map: dict[str, str], use_az: bool) -> dict[str, str]:
    """Resource name -> Azure scope: the map first, then (read-only) the Azure CLI for the rest."""
    found = {k.lower(): v for k, v in scope_map.items() if SCOPE.match(v)}
    if use_az:
        for n in sorted(x for x in names if x.lower() not in found):
            q = ["az", "group", "list", "--query", f"[?name=='{n}'].id | [0]", "-o", "tsv"] if guess_type(n) == RESOURCE_TYPES["rg"] \
                else ["az", "resource", "list", "--name", n, "--query", "[0].id", "-o", "tsv"]
            try:
                v = subprocess.run(q, capture_output=True, text=True, timeout=60, check=False).stdout.strip()
            except (OSError, subprocess.TimeoutExpired):
                v = ""
            if SCOPE.match(v):
                found[n.lower()] = v
    return found


def apply_test_scope(rows: list[dict], scope: str) -> None:
    """Pilot test: every Azure role is given the one test resource group; repeats of a role there are merged, naming
    the real resources it stands for."""
    for r in rows:
        merged: dict[str, dict] = {}
        for g in r["azureRoles"]:
            key = g["roleName"].lower()
            if key in merged:
                merged[key]["resourceName"] = f"{merged[key]['resourceName']}, {g['resourceName']}"[:200]
            else:
                merged[key] = {**g, "scope": scope, "resourceType": "Microsoft.Resources/resourceGroups"}
        r["azureRoles"] = list(merged.values())
        if merged:
            note = f"TEST: the roles are assigned in the test resource group {scope.rsplit('/', 1)[-1]} (stands for the real resources named on each line)."
            r["notes"] = f"{note} {r['notes'] or ''}".strip()[:2000]


def build(tenant: str, project: str, classification: Path | None, responses: list[Path], scope_map: dict[str, str], use_az: bool,
          resource_prefix: str = "", test_scope: str | None = None) -> dict:
    import openpyxl
    rbac = read_response(responses) if responses else {}
    rows: list[dict] = []
    envs: list[str] = []
    empty_envs: set[str] = set()
    listed: set[str] = set()
    domains: dict[str, str] = {}
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")

    if classification:
        wb = openpyxl.load_workbook(classification, data_only=False)
        for ws in wb.worksheets:
            m = re.match(rf"^{re.escape(tenant)}\s*-\s*(Dev|Test|PreProd|Prod)$", ws.title, re.I)
            if not m:
                continue  # tenant separation: other tenants' tabs are never read
            env = {"dev": "Dev", "test": "Test", "preprod": "PreProd", "prod": "Prod"}[m.group(1).lower()]
            envs.append(env)
            header_row = next((r for r in ws.iter_rows(min_row=1, max_row=8) if text(r[2].value).lower() == "group name"), None)
            if header_row is None:
                continue
            head = [text(c.value).lower() for c in header_row]
            idx = {k: next((i for i, h in enumerate(head) if h.startswith(k)), None)
                   for k in ("group name", "object id", "role / access", "intended job", "scope", "source", "notes")}
            count = 0
            for r in ws.iter_rows(min_row=header_row[0].row + 1):
                name = text(r[idx["group name"]].value)
                g = GROUP.match(name)
                if not g:
                    continue
                count += 1
                listed.add(name.lower())
                scope_text = text(r[idx["scope"]].value) if idx["scope"] is not None else ""
                dm = re.search(r"Domain:\s*(.+)", scope_text)
                domain_label = clean_domain(dm.group(1), g["domain"]) if dm else g["domain"]
                domains[g["domain"]] = domain_label
                grants = [dict(x) for x in rbac.get(name.lower(), [])]
                if not grants:  # the classification's own "role → resource" bullets
                    for line in scope_text.splitlines():
                        b = BULLET.match(line)
                        if b:
                            for res in [x.strip() for x in b["res"].split("/") if x.strip()]:
                                grants.append({"roleName": b["role"].strip(), "resourceName": res, "resourceType": guess_type(res)})
                role_text = text(r[idx["role / access"]].value).splitlines()
                oid = text(r[idx["object id"]].value)
                rows.append(row_doc(project, tenant, env, name, g, domain_label, role_text[0] if role_text else level_of(g["role"]),
                                    grants, flag_of(r[idx["group name"]]), text(r[idx["notes"]].value) if idx["notes"] is not None else "",
                                    text(r[idx["intended job"]].value) if idx["intended job"] is not None else "",
                                    oid if GUID.match(oid) else None, "Exists",
                                    f"{classification.name} / {ws.title}", now, resource_prefix))
            if count == 0:
                empty_envs.add(env)

    # Groups the response document names that the classification does not list (yet).
    for name, grants in rbac.items():
        g = GROUP.match(name)
        if not g or name in listed or not name.startswith(f"azr-{tenant.lower()[:3]}"):
            continue
        env = ENVIRONMENTS[g["tier"]]
        if env not in envs:
            envs.append(env)
        red = env in empty_envs or (classification is not None and env not in envs[:-1] and env == "Prod")
        note = (f"The classification lists no {env} groups for {tenant}: the group must be created (IaC) before anyone is added."
                if red else "Not in the live group classification: the group is created after approval (by the agent or the ID team, or IaC where tenant policy requires).")
        rows.append(row_doc(project, tenant, env, name, g, domains.get(g["domain"], g["domain"]), level_of(g["role"]), [dict(x) for x in grants],
                            "Red" if red else "Amber", note, "", None, "ToBeCreated",
                            ", ".join(p.name for p in responses), now, resource_prefix))

    names = {x["resourceName"] for r in rows for x in r["azureRoles"]}
    scopes = {} if test_scope else resolve(names, scope_map, use_az)
    for r in rows:
        for x in r["azureRoles"]:
            x["scope"] = scopes.get(x["resourceName"].lower())
            x["roleDefinitionId"] = None
    if test_scope:
        apply_test_scope(rows, test_scope)
    order = [e for e in ("Dev", "Test", "PreProd", "Prod") if e in envs or e == "Prod"]
    return {"project": {"projectId": project, "tenantName": tenant, "environments": order}, "mappingSheet": rows}


def row_doc(project, tenant, env, name, g, domain_label, role, grants, flag, notes, jobs, oid, state, source, now, prefix) -> dict:
    row_id = re.sub(r"[^A-Za-z0-9_.:-]", "-", f"{g['domain']}-{g['role']}-{env}".upper())[:64]
    privileged = g["role"].lower() == "admins" or has_word(role, PRIVILEGED_WORDS)
    return {
        "id": row_id, "projectId": project, "rowId": row_id, "resource": f"{prefix}{tenant} {domain_label}".strip(),
        "accessLevel": level_of(g["role"]), "environment": env, "entraGroup": name, "groupObjectId": oid, "role": role[:200],
        "roleScope": None, "azureRoles": grants, "jobRoles": jobs.replace("\n", "; ")[:500] or None, "flag": flag,
        "notes": notes[:2000] or None, "privileged": privileged, "sensitiveResource": "pii" in g["role"].lower(), "deprecated": False,
        "groupState": state, "sourceReference": source[:200], "version": 1, "updatedBy": "import_workbook", "updatedAt": now,
        "isSample": False, "schemaVersion": 2,
    }


ROLES = {"Requester", "Owner", "Security", "IdTeam", "RoleTeam", "PMO", "Operator"}
SAFE_ID = re.compile(r"^[A-Za-z0-9_.:\-]{1,128}$")
MAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
CONTACT_LISTS = ("pmo", "security", "idTeam", "roleTeam", "operators")
CONTACT_ONE = ("securityManager", "idTeamManager", "roleTeamManager", "ownerManagerFallback")


SIGN_IN_PREFIX = "signin:"  # same as AccessService.SignInPrefix


def user_key(user: str) -> str:
    """The userAccess user id: the app's user id (/.auth/me userId) as is, or a sign-in name (mail form, e.g.
    jane@contoso.com) as "signin:<lower case>" — the API matches it like the named Owner (guest #EXT# and live.com# forms
    of that mail count as the same person)."""
    u = user.strip()
    if "@" in u:
        if "#" in u or not MAIL.match(u):
            raise SystemExit(f"{user}: give the sign-in name as a plain mail address (e.g. jane@contoso.com), not the #EXT# or live.com# form")
        return SIGN_IN_PREFIX + u.lower()
    if not SAFE_ID.match(u) or u.lower().startswith(SIGN_IN_PREFIX):
        raise SystemExit(f"{user}: not a user id (letters, digits, . _ : -) or a sign-in name")
    return u


def user_doc(user: str, project: str, roles: list[str]) -> dict:
    if not roles or set(roles) - ROLES:
        raise SystemExit(f"roles must be from {sorted(ROLES)}")
    key = user_key(user)
    return {"id": f"{key}:{project}", "userId": key, "projectId": project, "roles": roles, "isSample": False, "schemaVersion": 1}


def users(pairs: list[str], project: str) -> list[dict]:
    """--register-user USER=Role,Role (USER = sign-in name or user id) -> userAccess documents (not sample data)."""
    out = []
    for pair in pairs:
        user, _, roles = pair.partition("=")
        wanted = [r.strip() for r in roles.split(",") if r.strip()]
        if not user or not wanted:
            raise SystemExit(f"--register-user must be USER=Role,Role (USER = sign-in name or user id) with roles from {sorted(ROLES)}")
        out.append(user_doc(user, project, wanted))
    return out


def project_doc(data: dict, live: dict | None, project_name: str, domain_suffix: str | None, tenant_id: str | None, test_contact: str | None) -> dict:
    """Pure: the project document after the import (keeps what was set live unless given again)."""
    p = data["project"]
    doc = dict(live) if live else {"id": p["projectId"], "projectId": p["projectId"], "contacts": {}, "schemaVersion": 2}
    doc.update({"name": project_name or doc.get("name") or f"{p['tenantName']} data platform", "tenantName": p["tenantName"],
                "environments": p["environments"], "isSample": False,
                "groupNamePattern": rf"^azr-{p['tenantName'].lower()[:3]}-rg-data-(dev|tst|ppd|prd)-[a-z0-9]+-weu-[A-Za-z0-9-]+$"})
    if domain_suffix:
        doc["domainIdSuffix"] = domain_suffix.lower()
    if tenant_id:
        doc["tenantId"] = tenant_id.lower()
    if test_contact:
        c = dict(doc.get("contacts") or {})
        for k in CONTACT_LISTS:
            c[k] = [test_contact.lower()]
        for k in CONTACT_ONE:
            c[k] = test_contact.lower()
        doc["contacts"] = c
    return doc


def apply(data: dict, project_name: str, domain_suffix: str | None, tenant_id: str | None, update: bool, database: str,
          test_contact: str | None = None, register: list[dict] | None = None) -> None:
    cs = os.environ.get("COSMOS_CONNECTION_STRING")
    if not cs:
        raise SystemExit("COSMOS_CONNECTION_STRING is not set")
    from azure.cosmos import CosmosClient
    from azure.cosmos.exceptions import CosmosResourceExistsError, CosmosResourceNotFoundError
    db = CosmosClient.from_connection_string(cs).get_database_client(database)
    projects = db.get_container_client("projects")
    p = data["project"]
    try:
        live = projects.read_item(p["projectId"], partition_key=p["projectId"])
    except CosmosResourceNotFoundError:
        live = None
    doc = project_doc(data, live, project_name, domain_suffix, tenant_id, test_contact)
    projects.upsert_item(doc)
    sheet = db.get_container_client("mappingSheet")
    created = updated = kept = 0
    for row in data["mappingSheet"]:
        try:
            sheet.create_item(row)
            created += 1
        except CosmosResourceExistsError:
            if not update:
                kept += 1
                continue
            live_row = sheet.read_item(row["id"], partition_key=row["projectId"])
            sheet.replace_item(row["id"], {**row, "version": int(live_row.get("version", 1)) + 1})  # in-flight requests are reviewed again
            updated += 1
    access = db.get_container_client("userAccess")
    for u in register or []:
        access.upsert_item(u)
        print(f"  registered {u['userId']} on {u['projectId']}: {', '.join(u['roles'])}")
    contacts = "test contact " + test_contact if test_contact else ("contacts kept" if doc.get("contacts") else "NO contacts (no mail goes out)")
    print(f"project {p['projectId']} ({p['tenantName']}, {doc.get('tenantId') or 'agent tenant'}, {doc.get('domainIdSuffix') or 'no domain'}, {contacts}): "
          f"{created} rows created, {updated} updated, {kept} kept (use --update to replace)")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tenant", required=True, help="tenant name as in the workbook tabs, e.g. PEST or Vestacy")
    ap.add_argument("--project", required=True, help="project id (one project per tenant)")
    ap.add_argument("--project-name", default="")
    ap.add_argument("--classification", type=Path)
    ap.add_argument("--response", type=Path, action="append", default=[])
    ap.add_argument("--scope-map", type=Path)
    ap.add_argument("--resolve", action="store_true", help="resolve resource names with the Azure CLI (read-only)")
    ap.add_argument("--resource-prefix", default="", help="prefix for resource names (e.g. 'SAMPLE ' for a test copy)")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--update", action="store_true", help="with --apply: replace existing rows (bumps their version)")
    ap.add_argument("--domain-suffix")
    ap.add_argument("--tenant-id", help="Entra tenant GUID of this project (the agent's own tenant for a pilot test)")
    ap.add_argument("--test-scope", help="pilot test: every Azure role in this one resource group (/subscriptions/<id>/resourceGroups/<name>)")
    ap.add_argument("--test-contact", help="test environment: every project contact (PMO, Security, ID team, role team, operators) is this mailbox")
    ap.add_argument("--register-user", action="append", default=[], metavar="USER=Role,Role",
                    help="register a person on the project by sign-in name (jane@contoso.com) or user id (repeatable)")
    ap.add_argument("--database", default=os.environ.get("COSMOS_DATABASE", "pmoreview"))
    args = ap.parse_args()
    if not args.classification and not args.response:
        ap.error("give --classification and/or --response")
    if args.tenant_id and not GUID.match(args.tenant_id):
        ap.error("--tenant-id must be the tenant's GUID")
    if args.test_scope and not re.match(r"^/subscriptions/[0-9a-fA-F-]{36}/resourceGroups/[A-Za-z0-9._()-]{1,90}$", args.test_scope):
        ap.error("--test-scope must be /subscriptions/<id>/resourceGroups/<name>")
    if args.test_contact and not MAIL.match(args.test_contact):
        ap.error("--test-contact must be one mail address")
    if args.project.upper().startswith("SAMPLE-"):
        ap.error("real workbooks go into a real project id (e.g. PEST), not a SAMPLE- one")
    register = users(args.register_user, args.project)
    scope_map = json.loads(args.scope_map.read_text(encoding="utf-8")) if args.scope_map else {}
    data = build(args.tenant, args.project, args.classification, args.response, scope_map, args.resolve, args.resource_prefix, args.test_scope)
    out = args.out or Path(__file__).resolve().parent.parent / ".local" / f"import-{args.project.lower()}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    rows = data["mappingSheet"]
    unresolved = sum(1 for r in rows for g in r["azureRoles"] if not g["scope"])
    print(f"{len(rows)} rows for {args.tenant} ({', '.join(data['project']['environments'])}): "
          f"{sum(r['flag'] == 'Red' for r in rows)} red, {sum(r['flag'] == 'Amber' for r in rows)} amber, "
          f"{sum(r['groupState'] == 'ToBeCreated' for r in rows)} to be created, "
          f"{sum(len(r['azureRoles']) for r in rows)} Azure roles ({unresolved} without a resolved scope). Wrote {out}")
    if args.apply:
        apply(data, args.project_name, args.domain_suffix, args.tenant_id, args.update, args.database, args.test_contact, register)
    return 0


if __name__ == "__main__":
    sys.exit(main())
